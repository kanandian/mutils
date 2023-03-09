# import docx
from docx import Document
import xlwt
from collections import defaultdict
import re
import os

import extract_info.extract_funcs as ef
from extract_info.config import all_items, selected_items, separators, folder_path


item_func = {
    '案号': ef.handle_anhao,
    '被告人': ef.handle_beigaoren,
    '本院认为': ef.handle_benyuanrenwei
}

def write_to_excel(data):
    global item_cnt
    for idx, item in enumerate(selected_items):
        if type(data[item]) == list:
            cont = ','.join(data[item])
        else:
            cont = data[item]
        worksheet.write(item_cnt, idx, cont)
    item_cnt+=1


def handle_files(file_path):
    print(file_path)
    article = list()
    # data = defaultdict(list)
    data = dict()
    doc = Document(file_path)
    for idx, para in enumerate(doc.paragraphs):
        txt = para.text
        if txt is None or txt == '':
            continue
        article.append(txt)
    for item in selected_items:
        data[item]=item_func[item](article)
    write_to_excel(data)


item_cnt=1
separators = '|'.join(separators)
files = os.listdir(folder_path)
workbook = xlwt.Workbook(encoding='utf8')
worksheet = workbook.add_sheet('1')
for idx, item in enumerate(selected_items):
    worksheet.write(0, idx, item)
for file_name in files:
    if file_name.endswith('.docx'):
        file_path = folder_path+'/'+file_name
        handle_files(file_path)
workbook.save('output1.xls')